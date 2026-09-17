"""Tests de l'import/export Scodoc."""

import os

from openpyxl import Workbook, load_workbook

from qcm_papier import scodoc


def test_load_students_table(tmp_path):
    """Charge une table étudiants Scodoc (etudid, code_nip, nom, prenom).

    Le code_nip Scodoc est numérique (ex: 12504873) ; l'ID étudiant lu sur
    la copie est 'p' + nip sans son 1er chiffre (ex: p2504873), comme le
    code JS (id = 'p' + nip.substring(1)).
    """
    wb = Workbook()
    ws = wb.active
    ws.append(["etudid", "code_nip", "nom", "prenom"])
    ws.append(["12345", "p1000001", "DUPONT", "Jean"])
    ws.append(["12346", "p1000002", "MARTIN", "Marie"])
    path = tmp_path / "etudiants.xlsx"
    wb.save(str(path))

    students = scodoc.load_students_table(str(path))
    assert len(students) == 2
    assert "p1000001" in students
    s = students["p1000001"]
    assert s.eid == "12345"
    assert s.nip == "p1000001"
    assert s.name == "DUPONT"
    assert s.firstname == "Jean"


def test_load_students_table_nip_numerique(tmp_path):
    """code_nip numérique Scodoc (sans 'p' initial) → 'p' + nip[1:].

    Cas réel : nip=12504873 (8 chiffres) → id='p2504873' (7 chiffres),
    qui correspond à l'identifiant lu sur la copie scannée.
    """
    wb = Workbook()
    ws = wb.active
    ws.append(["etudid", "code_nip", "nom", "prenom"])
    ws.append(["1416", "12504873", "DUPONT", "Jean"])
    ws.append(["1417", "12505117", "MARTIN", "Marie"])
    path = tmp_path / "etudiants.xlsx"
    wb.save(str(path))

    students = scodoc.load_students_table(str(path))
    assert "p2504873" in students  # 12504873 → p2504873
    assert "p2505117" in students  # 12505117 → p2505117
    s = students["p2504873"]
    assert s.eid == "1416"
    assert s.nip == "12504873"


def test_load_students_table_colonnes_manquantes(tmp_path):
    """Lève une erreur si les colonnes attendues sont absentes."""
    wb = Workbook()
    ws = wb.active
    ws.append(["foo", "bar"])
    ws.append(["1", "2"])
    path = tmp_path / "bad.xlsx"
    wb.save(str(path))
    try:
        scodoc.load_students_table(str(path))
        raise AssertionError("Devrait lever une erreur")
    except ValueError as e:
        assert "etudid" in str(e)


def test_export_scodoc_notes(tmp_path):
    """Injecte des notes dans une feuille Scodoc et vérifie le résultat."""
    wb = Workbook()
    ws = wb.active
    for i in range(7):
        ws.append([f"header{i}"])
    ws.append(["!12345", "", "", "", ""])
    ws.append(["!12346", "", "", "", ""])
    ws.append(["!99999", "", "", "", ""])
    path_in = tmp_path / "notes_in.xlsx"
    wb.save(str(path_in))

    path_out = tmp_path / "notes_out.xls"
    notes = {"12345": 15.0, "12346": 8.5}
    count = scodoc.export_scodoc_notes(str(path_in), str(path_out), notes, note_max=20.0, notemax=20.0, header_row=7)
    assert count == 2
    # Le fichier .xls est en réalité du .xlsx (openpyxl ne gère pas biff8).
    out_xlsx = str(path_out)[:-4] + ".xlsx"
    assert os.path.exists(out_xlsx)
    wb2 = load_workbook(out_xlsx)
    ws2 = wb2.active
    # Note du 12345 (ligne 8, colonne E).
    assert ws2.cell(row=8, column=5).value == "15,00"
    assert ws2.cell(row=9, column=5).value == "8,50"
    # L'étudiant 99999 sans note → vide.
    assert ws2.cell(row=10, column=5).value is None


def test_export_scodoc_notes_min0(tmp_path):
    """Avec min0=True, les notes négatives sont ramenées à 0."""
    wb = Workbook()
    ws = wb.active
    for i in range(7):
        ws.append([f"header{i}"])
    ws.append(["!12345", "", "", "", ""])
    ws.append(["!12346", "", "", "", ""])
    ws.append(["!12347", "", "", "", ""])
    path_in = tmp_path / "notes_in.xlsx"
    wb.save(str(path_in))

    path_out = tmp_path / "notes_out.xlsx"
    # 12345 : note négative ; 12346 : note positive ; 12347 : zéro.
    notes = {"12345": -1.5, "12346": 8.0, "12347": 0.0}
    count = scodoc.export_scodoc_notes(str(path_in), str(path_out), notes, min0=True)
    assert count == 3
    wb2 = load_workbook(str(path_out))
    ws2 = wb2.active
    assert ws2.cell(row=8, column=5).value == "0,00"  # -1.5 → 0
    assert ws2.cell(row=9, column=5).value == "8,00"
    assert ws2.cell(row=10, column=5).value == "0,00"

    # Sans min0, la note négative est conservée.
    path_out2 = tmp_path / "notes_out_raw.xlsx"
    scodoc.export_scodoc_notes(str(path_in), str(path_out2), notes, min0=False)
    wb3 = load_workbook(str(path_out2))
    ws3 = wb3.active
    assert ws3.cell(row=8, column=5).value == "-1,50"  # négatif conservé


def test_find_page_note():
    """find_page_note retrouve la note d'un EID parmi une liste."""
    notes = [("12345", 15.0), ("12346", None), ("12347", 8.0)]
    assert scodoc.find_page_note("12345", notes) == 15.0
    assert scodoc.find_page_note("12347", notes) == 8.0
    assert scodoc.find_page_note("12346", notes) is None  # note None
    assert scodoc.find_page_note("99999", notes) is None  # absent
